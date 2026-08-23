#!/usr/bin/env python3
"""Build the four ScholarOne / supervisor Word files from the markdown sources."""

from __future__ import annotations

import re
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Pt, Inches, RGBColor

ROOT = Path(__file__).resolve().parents[1]
WORD = Path(__file__).resolve().parent
MANUSCRIPT = ROOT / "manuscript.md"
TITLE_PAGE = ROOT / "title-page.md"
COVER = ROOT / "cover-letter.md"


def make_reference_doc(path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.font.color.rgb = RGBColor(0, 0, 0)
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_after = Pt(0)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        if name in doc.styles:
            doc.styles[name].font.name = "Times New Roman"
            doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    doc.save(path)


def strip_working_copy(text: str) -> str:
    return re.sub(
        r"\*Working copy with names\.[^\n]*\n+",
        "",
        text,
        count=1,
    )


def anonymized_body(text: str) -> str:
    text = strip_working_copy(text)
    # Drop author block, acknowledgements, disclosure, funding, data availability.
    text = re.sub(
        r"\nHeena Rathore \(corresponding author\).*?\n## Abstract",
        "\n## Abstract",
        text,
        count=1,
        flags=re.S,
    )
    return text


def pandoc_md(md_text: str, out_docx: Path, reference: Path, extra_args: list[str] | None = None) -> None:
    with tempfile.NamedTemporaryFile("w", suffix=".md", delete=False, encoding="utf-8") as tmp:
        tmp.write(md_text)
        src = Path(tmp.name)
    cmd = [
        "pandoc",
        str(src),
        "-o",
        str(out_docx),
        "--from",
        "markdown",
        "--to",
        "docx",
        "--reference-doc",
        str(reference),
    ]
    if extra_args:
        cmd.extend(extra_args)
    subprocess.run(cmd, check=True)
    src.unlink(missing_ok=True)


def drop_duplicate_title(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    seen = None
    for para in doc.paragraphs:
        t = para.text.strip()
        if not t:
            continue
        if seen is None:
            seen = t
            continue
        if t == seen:
            para.clear()
            break
    doc.save(str(docx_path))


def main() -> None:
    named = strip_working_copy(MANUSCRIPT.read_text(encoding="utf-8"))
    anon = anonymized_body(MANUSCRIPT.read_text(encoding="utf-8"))
    title = TITLE_PAGE.read_text(encoding="utf-8")
    cover = COVER.read_text(encoding="utf-8")

    with tempfile.TemporaryDirectory() as td:
        ref = Path(td) / "reference.docx"
        make_reference_doc(ref)
        named_out = WORD / "Rathore-Cao-Xu-When-Chinese-class-moved-home.docx"
        anon_out = WORD / "Rathore-manuscript-anonymized-for-review.docx"
        title_out = WORD / "Rathore-title-page.docx"
        cover_out = WORD / "Rathore-cover-letter-LCC.docx"
        pandoc_md(named, named_out, ref)
        pandoc_md(anon, anon_out, ref)
        pandoc_md(title, title_out, ref)
        pandoc_md(cover, cover_out, ref)
        drop_duplicate_title(named_out)

    print("Wrote:")
    for p in (
        WORD / "Rathore-Cao-Xu-When-Chinese-class-moved-home.docx",
        WORD / "Rathore-manuscript-anonymized-for-review.docx",
        WORD / "Rathore-title-page.docx",
        WORD / "Rathore-cover-letter-LCC.docx",
    ):
        print(" ", p.name, p.stat().st_size, "bytes")


if __name__ == "__main__":
    main()
