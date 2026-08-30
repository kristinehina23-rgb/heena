#!/usr/bin/env python3
"""Build the freshman-share PPTX without embedding 等线/OTTO fonts."""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from export_layout import export as export_layout
from make_ppt import build as build_show_pptx
from render_slides import render_all

ROOT = Path(__file__).resolve().parent
CONTENT = json.loads((ROOT / "content.json").read_text(encoding="utf-8"))
OUTPUTS = ROOT / "outputs"
SLIDE_DIR = OUTPUTS / "slides"
PPTX_NAME = "喜娜_中大本科新生分享_丰富经历照片终版_2026_editable.pptx"

# Office-safe TrueType family. Do not use 等线 / DengXian:
# artifact-tool cannot decode its embedded OTTO/CFF scaler (0x4F54544F),
# and the EOT family "等线" does not match the SFNT name "DengXian".
FONT = "Microsoft YaHei"
FONT_SERIF = "Noto Serif CJK SC"

WIDE = Inches(13.333333)
HIGH = Inches(7.5)
GREEN = RGBColor(0x00, 0x56, 0x1F)
GREEN_DEEP = RGBColor(0x00, 0x22, 0x0E)
CREAM = RGBColor(0xF6, 0xF1, 0xE4)
GOLD = RGBColor(0xD2, 0x98, 0x65)
INK = RGBColor(0x22, 0x1F, 0x1B)
MUTED = RGBColor(0x5C, 0x56, 0x4C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAPER = RGBColor(0xFB, 0xFA, 0xF6)


def resolve_photos() -> dict:
    photos = CONTENT["photos"]
    resolved = {}
    for key, spec in photos.items():
        slot = ROOT / spec["slot"]
        fallback = (ROOT / spec["fallback"]).resolve()
        src = slot if slot.exists() else fallback
        item = dict(spec)
        item["resolved"] = str(src)
        item["resolvedAbs"] = str(src)
        item["usingFallback"] = not slot.exists()
        resolved[key] = item
    (ROOT / "resolved-photos.json").write_text(
        json.dumps(resolved, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return resolved


def convert_emblems() -> None:
    assets = ROOT / "assets"
    assets.mkdir(exist_ok=True)
    pairs = (
        (ROOT.parent / "poster/assets/校徽_反白.svg", assets / "emblem-white.png"),
        (ROOT.parent / "poster/assets/校徽_绿.svg", assets / "emblem-green.png"),
    )
    for src, dest in pairs:
        if dest.exists() and dest.stat().st_mtime >= src.stat().st_mtime:
            continue
        subprocess.run(
            ["rsvg-convert", "-w", "256", "-h", "256", str(src), "-o", str(dest)],
            check=True,
        )


def set_run_font(run, name: str, size: int, bold: bool, color: RGBColor) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = rPr.find(qn(tag))
        xml = f'<{tag} xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="{name}"/>'
        if node is None:
            rPr.append(parse_xml(xml))
        else:
            node.set("typeface", name)


def add_text(slide, name, left, top, width, height, text, *, size=20, bold=False, color=INK, font=FONT, align=PP_ALIGN.LEFT, anchor="t"):
    box = slide.shapes.add_textbox(left, top, width, height)
    box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf._txBody.bodyPr.set("anchor", anchor)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run_font(run, font, size, bold, color)
    return box


def add_rect(slide, name, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.name = name
    shape.line.fill.background()
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    return shape


def add_picture(slide, name, path, left, top, width, height):
    pic = slide.shapes.add_picture(str(path), left, top, width, height)
    pic.name = name
    return pic


def add_notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text or ""


def brand(slide, light: bool, right: str, page: int, total: int, meta: dict) -> None:
    emblem = ROOT / ("assets/emblem-green.png" if light else "assets/emblem-white.png")
    add_picture(slide, "校徽", emblem, Inches(0.38), Inches(0.22), Inches(0.42), Inches(0.42))
    color = GREEN if light else CREAM
    muted = MUTED if light else RGBColor(0xDC, 0xD6, 0xC6)
    add_text(slide, "校名", Inches(0.9), Inches(0.2), Inches(4), Inches(0.28), meta["university"], size=14, bold=True, color=color, font=FONT_SERIF)
    add_text(slide, "校名英文", Inches(0.9), Inches(0.44), Inches(5), Inches(0.22), meta["universityEn"].upper(), size=9, color=muted)
    add_text(slide, "页眉", Inches(8.2), Inches(0.28), Inches(4.7), Inches(0.3), right, size=12, color=GOLD if not light else GREEN, align=PP_ALIGN.RIGHT)
    add_text(slide, "校训", Inches(0.38), Inches(7.12), Inches(8), Inches(0.24), meta["motto"], size=10, color=muted)
    add_text(slide, "页码", Inches(11.2), Inches(7.12), Inches(1.7), Inches(0.24), f"{page:02d} / {total:02d}", size=10, color=muted, align=PP_ALIGN.RIGHT)


def photo_of(photos: dict, key: str) -> Path:
    return Path(photos[key]["resolvedAbs"])


def build_native(photos: dict) -> Path:
    prs = Presentation()
    prs.slide_width = WIDE
    prs.slide_height = HIGH
    blank = prs.slide_layouts[6]
    meta = CONTENT["meta"]
    slides = CONTENT["slides"]
    total = len(slides)

    for i, data in enumerate(slides, start=1):
        slide = prs.slides.add_slide(blank)
        kind = data["kind"]
        if kind in {"cover", "close"}:
            add_picture(slide, "背景", ROOT / data["background"], 0, 0, WIDE, HIGH)
            add_rect(slide, "遮罩", 0, 0, WIDE, HIGH, GREEN_DEEP)
            # soften overlay by stacking a second cream-tinted rect? keep deep green.
            brand(slide, False, data.get("kicker", meta["series"]), i, total, meta)
            add_rect(slide, "金线", Inches(0.6), Inches(4.75), Inches(0.85), Inches(0.035), GOLD)
            add_text(slide, "标题 1", Inches(0.55), Inches(4.9), Inches(12), Inches(1.3), data["title"], size=36, bold=True, color=CREAM, font=FONT_SERIF)
            add_text(slide, "副标题", Inches(0.55), Inches(6.25), Inches(12), Inches(0.4), data.get("subtitle") or data.get("body", ""), size=16, color=GOLD)
            add_text(slide, "讲者", Inches(0.55), Inches(6.65), Inches(12), Inches(0.35), data.get("speaker", ""), size=16, color=CREAM)
        elif kind == "quote":
            add_rect(slide, "底", 0, 0, WIDE, HIGH, PAPER)
            brand(slide, True, data["kicker"], i, total, meta)
            add_text(slide, "标题 1", Inches(0.6), Inches(1.2), Inches(12), Inches(0.8), data["title"], size=32, bold=True, color=GREEN, font=FONT_SERIF)
            add_text(slide, "导语", Inches(0.6), Inches(2.1), Inches(12.1), Inches(1.6), data["body"], size=16, color=INK)
            for n, point in enumerate(data["points"]):
                left = Inches(0.6 + n * 4.2)
                add_rect(slide, f"要点底{n+1}", left, Inches(3.85), Inches(4.0), Inches(2.55), WHITE)
                add_rect(slide, f"要点金{n+1}", left, Inches(3.85), Inches(0.07), Inches(2.55), GOLD)
                add_text(slide, f"要点号{n+1}", left + Inches(0.25), Inches(4.05), Inches(3.5), Inches(0.35), f"0{n+1}", size=14, color=GOLD, bold=True)
                add_text(slide, f"要点{n+1}", left + Inches(0.25), Inches(4.5), Inches(3.5), Inches(1.6), point, size=16, color=INK)
        elif kind == "map":
            add_rect(slide, "底", 0, 0, WIDE, HIGH, PAPER)
            brand(slide, True, data["kicker"], i, total, meta)
            add_text(slide, "标题 1", Inches(0.6), Inches(1.15), Inches(12), Inches(0.7), data["title"], size=30, bold=True, color=GREEN, font=FONT_SERIF)
            for n, card in enumerate(data["cards"]):
                left = Inches(0.55 + n * 4.25)
                add_picture(slide, f"路线图{n+1}", ROOT / card["photo"], left, Inches(2.1), Inches(4.05), Inches(4.7))
                add_rect(slide, f"路线遮罩{n+1}", left, Inches(5.4), Inches(4.05), Inches(1.4), GREEN_DEEP)
                add_text(slide, f"路线号{n+1}", left + Inches(0.2), Inches(5.5), Inches(3.6), Inches(0.3), card["num"], size=12, color=GOLD)
                add_text(slide, f"路线标题{n+1}", left + Inches(0.2), Inches(5.8), Inches(3.6), Inches(0.4), card["title"], size=18, bold=True, color=CREAM, font=FONT_SERIF)
                add_text(slide, f"路线说明{n+1}", left + Inches(0.2), Inches(6.25), Inches(3.6), Inches(0.35), card["line"], size=13, color=CREAM)
        elif kind == "experience":
            add_rect(slide, "底", 0, 0, WIDE, HIGH, PAPER)
            brand(slide, True, f"{meta['series']} · {meta['year']}", i, total, meta)
            add_text(slide, "标题 1", Inches(0.4), Inches(0.9), Inches(5.8), Inches(1.15), data["title"], size=22, bold=True, color=GREEN, font=FONT_SERIF)
            add_text(slide, "文本框 25", Inches(0.4), Inches(2.1), Inches(5.6), Inches(0.35), data["greenTitle"], size=13, color=GREEN)
            add_rect(slide, "统计1底", Inches(0.4), Inches(2.55), Inches(2.7), Inches(1.05), WHITE)
            add_rect(slide, "统计2底", Inches(3.25), Inches(2.55), Inches(2.7), Inches(1.05), WHITE)
            add_text(slide, "文本框 7", Inches(0.55), Inches(2.58), Inches(2.4), Inches(0.55), data["n1"], size=28, bold=True, color=GREEN, font=FONT_SERIF)
            add_text(slide, "文本框 10", Inches(0.55), Inches(3.15), Inches(2.4), Inches(0.35), data["l1"], size=12, color=MUTED)
            add_text(slide, "文本框 12", Inches(3.4), Inches(2.58), Inches(2.4), Inches(0.55), data["n2"], size=28, bold=True, color=GREEN, font=FONT_SERIF)
            add_text(slide, "文本框 13", Inches(3.4), Inches(3.15), Inches(2.4), Inches(0.35), data["l2"], size=12, color=MUTED)
            add_rect(slide, "引言底", Inches(0.4), Inches(3.8), Inches(5.55), Inches(2.85), GREEN)
            add_text(slide, "文本框 22", Inches(0.6), Inches(4.0), Inches(5.2), Inches(0.85), data["whiteTitle"], size=16, bold=True, color=CREAM, font=FONT_SERIF)
            add_text(slide, "文本框 21", Inches(0.6), Inches(4.9), Inches(5.2), Inches(1.55), data["body"], size=13, color=CREAM)
            add_picture(slide, "图片 29", photo_of(photos, data["image1"]), Inches(6.2), Inches(0.85), Inches(3.35), Inches(5.95))
            add_picture(slide, "图片 23", photo_of(photos, data["image2"]), Inches(9.7), Inches(0.85), Inches(3.25), Inches(5.95))
        elif kind == "gallery":
            add_rect(slide, "底", 0, 0, WIDE, HIGH, PAPER)
            brand(slide, True, data["kicker"], i, total, meta)
            add_text(slide, "标题 1", Inches(0.55), Inches(1.05), Inches(12.2), Inches(0.8), data["title"], size=24, bold=True, color=GREEN, font=FONT_SERIF)
            for n, item in enumerate(data["items"]):
                col, row = n % 4, n // 4
                left = Inches(0.45 + col * 3.22)
                top = Inches(2.05 + row * 2.45)
                add_picture(slide, f"相册{n+1}", ROOT / item["photo"], left, top, Inches(3.05), Inches(2.25))
        elif kind == "advice":
            add_rect(slide, "底", 0, 0, WIDE, HIGH, PAPER)
            brand(slide, True, data["kicker"], i, total, meta)
            add_text(slide, "标题 1", Inches(0.55), Inches(1.15), Inches(12.2), Inches(1.0), data["title"], size=28, bold=True, color=GREEN, font=FONT_SERIF)
            for n, item in enumerate(data["items"]):
                left = Inches(0.55 + n * 4.25)
                add_rect(slide, f"建议底{n+1}", left, Inches(2.4), Inches(4.05), Inches(4.3), WHITE)
                add_text(slide, f"建议号{n+1}", left + Inches(0.25), Inches(2.55), Inches(3.5), Inches(0.35), item["n"], size=14, color=GOLD, bold=True)
                add_text(slide, f"建议标题{n+1}", left + Inches(0.25), Inches(3.0), Inches(3.55), Inches(1.1), item["title"], size=18, bold=True, color=GREEN, font=FONT_SERIF)
                add_text(slide, f"建议正文{n+1}", left + Inches(0.25), Inches(4.2), Inches(3.55), Inches(2.2), item["body"], size=14, color=INK)
        add_notes(slide, data.get("notes", ""))

    # Do not embed fonts. Office will substitute Microsoft YaHei / 微软雅黑.
    out = OUTPUTS / PPTX_NAME
    OUTPUTS.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out


def main() -> int:
    convert_emblems()
    photos = resolve_photos()
    missing = [k for k, v in photos.items() if v["usingFallback"]]
    if missing:
        print("using campus atmosphere photos for:", ", ".join(missing), flush=True)
    render_all(photos, SLIDE_DIR)
    editable = build_native(photos)
    show = build_show_pptx()
    export_layout(editable)
    print(show)
    print(editable)
    print("fonts: Microsoft YaHei / Noto Serif CJK SC (not 等线/DengXian, not embedded OTTO)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
