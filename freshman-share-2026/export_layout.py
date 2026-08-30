#!/usr/bin/env python3
"""Export artifact-tool-style layout JSON and inspect NDJSON from the PPTX."""

from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parent
CONTENT = json.loads((ROOT / "content.json").read_text(encoding="utf-8"))
PPTX = ROOT / "outputs" / "喜娜_中大本科新生分享_丰富经历照片终版_2026.pptx"
OUT = ROOT / "outputs" / "slides"
INSPECT = ROOT / "outputs" / "喜娜_中大本科新生分享_丰富经历照片终版_2026.pptx.inspect.ndjson"

EMU_PER_PT = 12700


def pt(emu: int) -> float:
    return round(float(emu) / EMU_PER_PT, 2)


def shape_kind(shape) -> str:
    if shape.has_text_frame and shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
        if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            return "shape"
        return "textbox"
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return "image"
    return "shape"


def shape_text(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return "\n".join(p.text for p in shape.text_frame.paragraphs).strip()


def export(pptx_path: Path = PPTX) -> Path:
    prs = Presentation(pptx_path)
    OUT.mkdir(parents=True, exist_ok=True)
    rows = []
    for i, slide in enumerate(prs.slides, start=1):
        meta = CONTENT["slides"][i - 1]
        elements = []
        for shape in slide.shapes:
            frame = {
                "left": pt(shape.left),
                "top": pt(shape.top),
                "width": pt(shape.width),
                "height": pt(shape.height),
            }
            kind = shape_kind(shape)
            text = shape_text(shape)
            node = {
                "kind": kind,
                "name": shape.name,
                "path": f"/slide[{i}]/{shape.name}",
                "frame": frame,
            }
            if text:
                node["text"] = text
            if kind == "image":
                node["image"] = True
            elements.append(node)
            rows.append(
                {
                    "slide": i,
                    "id": meta["id"],
                    "kind": kind,
                    "name": shape.name,
                    "text": text[:500],
                    "frame": frame,
                }
            )
        notes = slide.notes_slide.notes_text_frame.text
        elements.append(
            {
                "kind": "notes",
                "name": "speakerNotes",
                "path": f"/slide[{i}]/speakerNotes",
                "text": notes,
                "visible": True,
            }
        )
        rows.append(
            {
                "slide": i,
                "id": meta["id"],
                "kind": "notes",
                "name": "speakerNotes",
                "text": notes[:500],
            }
        )
        layout = {
            "slide": i,
            "id": meta["id"],
            "kind": meta["kind"],
            "frame": {
                "left": 0,
                "top": 0,
                "width": pt(prs.slide_width),
                "height": pt(prs.slide_height),
            },
            "elements": elements,
        }
        dest = OUT / f"slide-{i:02d}.layout.json"
        dest.write_text(json.dumps(layout, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    INSPECT.write_text("\n".join(json.dumps(r, ensure_ascii=False) for r in rows) + "\n", encoding="utf-8")
    print(INSPECT)
    return INSPECT


if __name__ == "__main__":
    export()
