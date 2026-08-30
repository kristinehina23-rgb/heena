#!/usr/bin/env python3
"""Structural checks for the rebuilt deck."""

from __future__ import annotations

import zipfile
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parent
PPTX = ROOT / "outputs" / "喜娜_中大本科新生分享_丰富经历照片终版_2026.pptx"
REQUIRED = {
    4: ["标题 1", "文本框 25", "文本框 7", "文本框 10", "文本框 12", "文本框 13", "文本框 22", "文本框 21", "图片 29", "图片 23"],
    5: ["标题 1", "文本框 25", "文本框 7", "文本框 10", "文本框 12", "文本框 13", "文本框 22", "文本框 21", "图片 29", "图片 23"],
    6: ["标题 1", "文本框 25", "文本框 7", "文本框 10", "文本框 12", "文本框 13", "文本框 22", "文本框 21", "图片 29", "图片 23"],
}


def main() -> None:
    assert PPTX.exists(), PPTX
    prs = Presentation(PPTX)
    assert len(prs.slides) == 8
    for idx, names in REQUIRED.items():
        have = {s.name for s in prs.slides[idx - 1].shapes}
        missing = [n for n in names if n not in have]
        assert not missing, (idx, missing)
        notes = prs.slides[idx - 1].notes_slide.notes_text_frame.text
        assert notes.strip(), idx

    with zipfile.ZipFile(PPTX) as z:
        assert not any("font" in n.lower() and n.endswith((".fntdata", ".odttf", ".eot")) for n in z.namelist())
        xml = b"".join(z.read(n) for n in z.namelist() if n.endswith(".xml"))
    assert b"DengXian" not in xml
    assert "等线".encode("utf-8") not in xml
    assert b"Microsoft YaHei" in xml
    print("ok")


if __name__ == "__main__":
    main()
